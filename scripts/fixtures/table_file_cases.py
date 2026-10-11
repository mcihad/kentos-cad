#!/usr/bin/env python3
"""The shared cases of the table file reader (docs/adr/0184 §4, 0223 §1).

    python3 scripts/fixtures/table_file_cases.py           # writes the files and the cases
    python3 scripts/fixtures/table_file_cases.py --check   # writes nothing; compares, then cross-checks

Writes the sample files of fixtures/table/v1/files (text tables in several
encodings and separators, Excel workbooks and an OpenDocument spreadsheet
built here with zipfile and hand-written XML, an old Excel file's head,
another archive, an OpenDocument text) and fixtures/table/v1/files.json,
each file's sheets as the rule reads them; libre.xlsx and libre.ods are
LibreOffice's own writing of two of them, kept as they are. The rule is
written here from the ADRs on their own, not from an implementation's
output: the workbooks are read with Python's zipfile and xml.etree, the
text with this script's own record reader. The formats core
(`table_file::read_with`, natively and through WASM: `readTableFile`) is
held to it. --check then reads every workbook with openpyxl (an
independent reader) and every OpenDocument spreadsheet through LibreOffice
(converted to XLSX, read with openpyxl), and compares their cells with the
rule's; either is skipped when it is not installed.

- A file over 64 MB is refused; one starting with D0 CF 11 E0 (an old Excel
  file) is refused; a zip archive is an Excel workbook when it holds
  xl/workbook.xml, an OpenDocument spreadsheet when it holds content.xml
  and its mimetype (if any) is a spreadsheet's, else refused.
- A text file's encoding: a byte order mark (EF BB BF UTF-8, FF FE UTF-16LE,
  FE FF UTF-16BE), else UTF-8 when the bytes are UTF-8, else Windows-1254.
- Its records: fields split by the separator; a field starting with a quote
  (spaces before it aside) is quoted up to the next lone quote ("" one
  quote), line breaks inside it kept, what follows it up to the separator
  added as it is; an unquoted field loses the spaces round it; a record ends
  at a line break (CR LF, LF or CR) outside quotes; a line with nothing on
  it is no record.
- The separator: of tab, semicolon, comma and bar, in that order, the first
  that gives the most of the first 50 records the field count most of them
  have (the larger of two counts as common), when that count is two or
  more; else runs of spaces and tabs (no quoting) when they give at least
  half of those records the count most of them have, two or more; else the
  whole line is one field (without the spaces round it).
- At most `rows` rows and `columns` columns are read (Tablo ekle: 10 001
  and 101; a case may say others). A text file with more records says it
  was cut. A workbook's sheet says it was cut when a filled cell lies in a
  row past the most, and its reading stops there (later rows are not read);
  a filled cell past the columns is left out.
- An Excel workbook's sheets in xl/workbook.xml's order, each by its name
  through its relationship (a target relative to xl/, or absolute from the
  root). A cell where its reference (A1, from 0) puts it, else after the one
  before in its row (a row where its r says, else after the one before).
  Its words: a shared string's text (its t elements' text, those under rPh
  left out), an inline string's alike, a formula's string result (t="str")
  as written, each with `_xHHHH_` read as the character it names (left to
  right, so `_x005F_x0041_` stays `_x0041_`); a boolean 1 DOĞRU, 0 YANLIŞ;
  a number (no t, or t="n") whose cell style's number format is a date
  format, its date; anything else its v's text as written; an empty value
  no cell.
- A date format (openpyxl's rule): its first section, without quoted text
  and without bracketed parts other than [h], [hh], [m], [mm], [s], [ss],
  holds d, m, h, y or s (either case) not straight after _ or \\; but a
  format starting with [h], [hh], [m], [mm], [s] or [ss] (either case) is
  an elapsed time and stays a number. d or y, or m without h and s, make a
  date; h or s a time; both a date and time. The built-in formats are
  Excel's (14–22, 45–47); the workbook's numFmts name the others.
- A date's text: the number v (0 or more, up to 2 958 466) split into its
  whole day and its fraction, the fraction's seconds rounded half up (86 400
  of them the next day, but for a date alone); a time, or any v below 1 on
  day 0, HH:MM:SS; a date YYYY-MM-DD, a date and time YYYY-MM-DDTHH:MM:SS;
  the day counted from 30 December 1899 (one more below 60, but not at 0)
  or, for a workbook with date1904, from 1 January 1904. Otherwise the
  number as written.
- An OpenDocument spreadsheet's tables (content.xml's table:table elements
  not inside another) in order, by name. Its rows: table-row elements met
  under the table (in header rows, row groups and row lists too, not in
  sub-tables), each repeated table:number-rows-repeated times (a count of 1
  or more, else 1); an empty one only moves the row on. Their cells:
  table-cell and covered-table-cell (always empty), each repeated
  table:number-columns-repeated times. A cell's words by office:value-type:
  float, percentage, currency office:value; date office:date-value (a
  date and time with a fraction of a second, YYYY-MM-DDTHH:MM:SS.f, rounded
  half up to the second, into the next day if so); time
  office:time-value read as PT<h>H<m>M<s>S into HH:MM:SS (seconds rounded
  half up, hours past 24 kept; as written when not so); boolean true DOĞRU,
  false YANLIŞ, else as written; other types office:string-value; when the
  attribute is not there, the text of the cell's text:p and text:h children
  joined by line breaks: their text as written (LibreOffice Calc does not
  fold a cell's white space as ODF does a document's), a text:s c spaces (at
  most 10 000), a text:tab a tab, a text:line-break a line break, other text
  elements by their text, text:note and text:ruby-text and elements of other
  namespaces left out (the text after them kept).
"""
import io
import json
import math
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DIR = ROOT / "fixtures/table/v1"
FILES = DIR / "files"
OUT = DIR / "files.json"
MAX_BYTES = 64 * 1024 * 1024
ROWS, COLUMNS = 10_001, 101
PROCESSING = {"rows": 250_001, "columns": 501}
SAMPLE = 50

LABELS = {"utf-8": "UTF-8", "utf-16-le": "UTF-16", "utf-16-be": "UTF-16", "cp1254": "Windows-1254 (Türkçe)"}
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
OFFICE = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"
TABLE = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
TEXTNS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"

# ── Text ────────────────────────────────────────────────────────────────


def encoding(data):
    if data.startswith(b"\xef\xbb\xbf"):
        return "utf-8", 3
    if data.startswith(b"\xff\xfe"):
        return "utf-16-le", 2
    if data.startswith(b"\xfe\xff"):
        return "utf-16-be", 2
    try:
        data.decode("utf-8")
        return "utf-8", 0
    except UnicodeDecodeError:
        return "cp1254", 0


def records(text, sep, limit):
    """Records split by sep (a character, 'spaces' or 'line'), at most limit; and whether more followed."""
    out = []
    i, n = 0, len(text)
    while i < n:
        if len(out) == limit:
            return out, True
        fields, field, quoted, was_quoted, anything = [], "", False, False, False
        while i < n:
            c = text[i]
            i += 1
            if quoted:
                if c == '"':
                    if i < n and text[i] == '"':
                        field += '"'
                        i += 1
                    else:
                        quoted = False
                else:
                    field += c
                continue
            if c in "\r\n":
                if c == "\r" and i < n and text[i] == "\n":
                    i += 1
                break
            anything = True
            if sep not in ("spaces", "line") and c == sep:
                fields.append(field if was_quoted else field.strip(" "))
                field, was_quoted = "", False
            elif sep not in ("spaces", "line") and c == '"' and field.strip(" ") == "" and not was_quoted:
                field, quoted, was_quoted = "", True, True
            elif sep == "spaces" and c in " \t":
                if field:
                    fields.append(field)
                    field = ""
            else:
                field += c
        if not anything and field == "" and not fields and not was_quoted:
            continue
        if sep == "spaces":
            if field:
                fields.append(field)
        else:
            fields.append(field if was_quoted else field.strip(" "))
        out.append(fields)
    return out, False


def steadiness(rows):
    counts = {}
    for r in rows:
        counts[len(r)] = counts.get(len(r), 0) + 1
    if not counts:
        return 0, 0
    return max((k, n) for n, k in counts.items())


def separator(text):
    best = None
    for c in ["\t", ";", ",", "|"]:
        sample, _ = records(text, c, SAMPLE)
        alike, count = steadiness(sample)
        if count >= 2 and (best is None or alike > best[0]):
            best = (alike, c)
    if best:
        return best[1]
    sample, _ = records(text, "spaces", SAMPLE)
    alike, count = steadiness(sample)
    return "spaces" if count >= 2 and 2 * alike >= len(sample) else "line"


def read_text(data, rows=ROWS, columns=COLUMNS):
    enc, skip = encoding(data)
    text = data[skip:].decode(enc, errors="replace")
    out, cut = records(text, separator(text), rows)
    out = [r[:columns] for r in out]
    return {"sheets": [{"rows": out, **({"cut": True} if cut else {})}], "encoding": LABELS[enc]}


class Grid:
    """A sheet's rows as cells come: a filled cell past the rows cuts the sheet (put says stop)."""

    def __init__(self, rows, columns):
        self.rows, self.columns, self.out, self.cut = rows, columns, [], False

    def put(self, i, j, words):
        if words == "":
            return True
        if i >= self.rows:
            self.cut = True
            return False
        if j >= self.columns:
            return True
        while len(self.out) <= i:
            self.out.append([])
        while len(self.out[i]) <= j:
            self.out[i].append("")
        self.out[i][j] = words
        return True

    def sheet(self, name):
        return {"name": name, "rows": self.out, **({"cut": True} if self.cut else {})}


# ── Excel workbooks ─────────────────────────────────────────────────────


def unescape(text):
    out, i = [], 0
    while i < len(text):
        h = text[i + 2 : i + 6]
        if (
            text.startswith("_x", i)
            and len(h) == 4
            and all(c in "0123456789abcdefABCDEF" for c in h)
            and text[i + 6 : i + 7] == "_"
            and not 0xD800 <= int(h, 16) <= 0xDFFF
        ):
            out.append(chr(int(h, 16)))
            i += 7
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def runs_text(node):
    out = []

    def walk(n, phonetic):
        tag = n.tag.split("}")[-1]
        phonetic = phonetic or tag == "rPh"
        if tag == "t" and not phonetic:
            out.append(n.text or "")
        for ch in n:
            walk(ch, phonetic)

    walk(node, False)
    return "".join(out)


def reference(r):
    letters = ""
    for ch in r:
        if ch.isascii() and ch.isalpha():
            letters += ch
        else:
            break
    if not letters or len(letters) > 3:
        return None
    col = 0
    for ch in letters.upper():
        col = col * 26 + (ord(ch) - ord("A") + 1)
    digits = r[len(letters):]
    if not digits.isdigit():
        return None
    row = int(digits)
    if row < 1:
        return None
    return col - 1, row - 1


BUILTIN = {14: "mm-dd-yy", 15: "d-mmm-yy", 16: "d-mmm", 17: "mmm-yy", 18: "h:mm AM/PM", 19: "h:mm:ss AM/PM",
           20: "h:mm", 21: "h:mm:ss", 22: "m/d/yy h:mm", 45: "mm:ss", 46: "[h]:mm:ss", 47: "mmss.0"}


def clock(fmt):
    """What a number format shows of a date: 'date', 'time', 'datetime', or None."""
    first = fmt.split(";")[0]
    if any(first.lower().startswith(e) for e in ("[h]", "[hh]", "[m]", "[mm]", "[s]", "[ss]")):
        return None
    kept, i = "", 0
    while i < len(first):
        c = first[i]
        if c == '"':
            j = first.find('"', i + 1)
            if j < 0:
                kept += first[i:]
                break
            i = j + 1
        elif c == "[":
            j = first.find("]", i)
            if j < 0:
                kept += first[i:]
                break
            if first[i + 1 : j] in ("h", "hh", "m", "mm", "s", "ss"):
                kept += first[i : j + 1]
            i = j + 1
        else:
            kept += c
            i += 1
    has = {"date": False, "time": False, "m": False}
    for k, c in enumerate(kept):
        if k > 0 and kept[k - 1] in "_\\":
            continue
        c = c.lower()
        if c in "dy":
            has["date"] = True
        elif c in "hs":
            has["time"] = True
        elif c == "m":
            has["m"] = True
    is_date = has["date"] or (has["m"] and not has["time"])
    if is_date and has["time"]:
        return "datetime"
    if is_date:
        return "date"
    return "time" if has["time"] else None


def date_text(v, kind, date1904):
    try:
        x = float(v.strip())
    except ValueError:
        return None
    if not math.isfinite(x) or x < 0 or x > 2_958_466:
        return None
    day = math.floor(x)
    secs = math.floor((x - day) * 86_400 + 0.5)
    if secs >= 86_400:
        secs -= 86_400
        if kind != "date":
            day += 1
    hms = f"{secs // 3600:02}:{secs // 60 % 60:02}:{secs % 60:02}"
    if kind == "time" or (x < 1 and day == 0):
        return hms
    if date1904:
        d = date(1904, 1, 1).toordinal() + day
    else:
        d = date(1899, 12, 30).toordinal() + day + (1 if 0 < x < 60 else 0)
    if d > date(9999, 12, 31).toordinal():
        return None
    iso = date.fromordinal(d).isoformat()
    return iso if kind == "date" else f"{iso}T{hms}"


def read_book(z, rows, columns):
    names = z.namelist()
    book = ET.fromstring(z.read("xl/workbook.xml"))
    pr = book.find("m:workbookPr", NS)
    date1904 = pr is not None and pr.get("date1904") in ("1", "true")
    targets = {}
    if "xl/_rels/workbook.xml.rels" in names:
        for rel in ET.fromstring(z.read("xl/_rels/workbook.xml.rels")):
            targets[rel.get("Id")] = rel.get("Target")
    shared = []
    if "xl/sharedStrings.xml" in names:
        shared = [unescape(runs_text(si)) for si in ET.fromstring(z.read("xl/sharedStrings.xml")) if si.tag.endswith("}si")]
    kinds = []
    if "xl/styles.xml" in names:
        styles = ET.fromstring(z.read("xl/styles.xml"))
        custom = {int(f.get("numFmtId")): f.get("formatCode") for f in styles.iter("{%s}numFmt" % NS["m"])}
        xfs = styles.find("m:cellXfs", NS)
        for xf in xfs.findall("m:xf", NS) if xfs is not None else []:
            fid = xf.get("numFmtId")
            fmt = custom.get(int(fid), BUILTIN.get(int(fid))) if fid and fid.isdigit() else None
            kinds.append(clock(fmt) if fmt is not None else None)
    sheets = []
    for sheet in book.iter("{%s}sheet" % NS["m"]):
        target = targets.get(sheet.get(REL))
        if target is None:
            continue
        path = target[1:] if target.startswith("/") else "xl/" + target
        if path not in names:
            continue
        doc = ET.fromstring(z.read(path))
        grid, next_row, stop = Grid(rows, columns), 0, False
        for row in doc.iter("{%s}row" % NS["m"]):
            r = row.get("r")
            i = int(r) - 1 if r and r.isdigit() and int(r) >= 1 else next_row
            next_row = i + 1
            next_col = 0
            for c in row.findall("m:c", NS):
                ref = reference(c.get("r")) if c.get("r") else None
                j = ref[0] if ref else next_col
                next_col = j + 1
                v = c.find("m:v", NS)
                v = v.text if v is not None and v.text is not None else ""
                t = c.get("t")
                if t == "s":
                    try:
                        k = int(v.strip())
                        words = shared[k] if 0 <= k < len(shared) else ""
                    except ValueError:
                        words = ""
                elif t == "inlineStr":
                    inline = c.find("m:is", NS)
                    words = unescape(runs_text(inline)) if inline is not None else ""
                elif t == "str":
                    words = unescape(v)
                elif t == "b":
                    words = {"1": "DOĞRU", "0": "YANLIŞ"}.get(v.strip(), v)
                elif t in (None, "n"):
                    s = c.get("s")
                    k = int(s) if s and s.isdigit() else 0
                    kind = kinds[k] if k < len(kinds) else None
                    words = (date_text(v, kind, date1904) if kind else None) or v
                else:
                    words = v
                if not grid.put(i, j, words):
                    stop = True
                    break
            if stop:
                break
        sheets.append(grid.sheet(sheet.get("name")))
    if not sheets:
        return {"sheets": [], "problem": "Çalışma kitabında okunacak sayfa yok."}
    return {"sheets": sheets}


# ── OpenDocument spreadsheets ──────────────────────────────────────────


def q(ns, local):
    return "{%s}%s" % (ns, local)


def repeat(el, name):
    v = (el.get(q(TABLE, name)) or "").strip()
    return int(v) if v.isdigit() and int(v) > 0 else 1


def ods_clock(t):
    s = t.strip()
    try:
        if not s.startswith("PT"):
            raise ValueError
        h, rest = s[2:].split("H", 1)
        m, rest = rest.split("M", 1)
        if not rest.endswith("S"):
            raise ValueError
        sec = float(rest[:-1])
        if not h.isdigit() or not m.isdigit() or not math.isfinite(sec) or sec < 0:
            raise ValueError
    except ValueError:
        return t
    total = int(h) * 3600 + int(m) * 60 + math.floor(sec + 0.5)
    return f"{total // 3600:02}:{total // 60 % 60:02}:{total % 60:02}"


def paragraph_text(p):
    """A paragraph's text, its white space as written (as LibreOffice Calc reads a cell; ODF's folding of white
    space in documents is not Calc's): text:s c spaces, text:tab a tab, text:line-break a line break."""
    out = [p.text or ""]
    for ch in p:
        if ch.tag == q(TEXTNS, "s"):
            c = (ch.get(q(TEXTNS, "c")) or "").strip()
            out.append(" " * min(int(c) if c.isdigit() else 1, 10_000))
        elif ch.tag == q(TEXTNS, "tab"):
            out.append("\t")
        elif ch.tag == q(TEXTNS, "line-break"):
            out.append("\n")
        elif ch.tag in (q(TEXTNS, "note"), q(TEXTNS, "ruby-text")) or not ch.tag.startswith("{%s}" % TEXTNS):
            pass
        else:
            out.append(paragraph_text(ch))
        out.append(ch.tail or "")
    return "".join(out)


def ods_date(v):
    t = v.strip()
    if len(t) > 20 and t[19] == "." and t[20:].isdigit() and t[20:].isascii():
        try:
            dt = datetime.strptime(t[:19], "%Y-%m-%dT%H:%M:%S")
        except ValueError:
            return v
        if not t[:19].isascii() or not t[0:4].isdigit():
            return v
        if t[20] >= "5":
            from datetime import timedelta

            dt += timedelta(seconds=1)
        return dt.isoformat()
    return v


def cell_words(c):
    kind = c.get(q(OFFICE, "value-type"))
    attr = {"float": "value", "percentage": "value", "currency": "value", "date": "date-value",
            "time": "time-value", "boolean": "boolean-value"}.get(kind, "string-value")
    v = c.get(q(OFFICE, attr))
    if v is not None:
        if kind == "time":
            return ods_clock(v)
        if kind == "date":
            return ods_date(v)
        if kind == "boolean":
            return {"true": "DOĞRU", "false": "YANLIŞ"}.get(v.strip(), v)
        return v
    lines = [paragraph_text(p) for p in c if p.tag in (q(TEXTNS, "p"), q(TEXTNS, "h"))]
    return "\n".join(lines)


def ods_rows(table):
    """The table's rows in order: (repeat, [cells as (repeat, words)])."""
    for el in table:
        if el.tag == q(TABLE, "table-row"):
            cells = []
            for c in el:
                if c.tag == q(TABLE, "covered-table-cell"):
                    cells.append((repeat(c, "number-columns-repeated"), ""))
                elif c.tag == q(TABLE, "table-cell"):
                    cells.append((repeat(c, "number-columns-repeated"), cell_words(c)))
            yield repeat(el, "number-rows-repeated"), cells
        elif el.tag == q(TABLE, "table"):
            continue
        else:
            yield from ods_rows(el)


def tables_of(el):
    for ch in el:
        if ch.tag == q(TABLE, "table"):
            yield ch
        else:
            yield from tables_of(ch)


def read_ods(z, rows, columns):
    doc = ET.fromstring(z.read("content.xml"))
    sheets = []
    for t in tables_of(doc):
        grid, row, stop = Grid(rows, columns), 0, False
        for n, cells in ods_rows(t):
            placed, col = [], 0
            for k, words in cells:
                if words and col < columns:
                    placed += [(j, words) for j in range(col, min(col + k, columns))]
                col += k
            if placed and not stop:
                for i in range(row, row + n):
                    if not all(grid.put(i, j, w) for j, w in placed):
                        stop = True
                        break
            row += n
            if stop:
                break
        sheets.append(grid.sheet(t.get(q(TABLE, "name")) or ""))
    if not sheets:
        return {"sheets": [], "problem": "Hesap tablosunda okunacak sayfa yok."}
    return {"sheets": sheets}


def read_archive(data, rows, columns):
    z = zipfile.ZipFile(io.BytesIO(data))
    names = z.namelist()
    if "xl/workbook.xml" in names:
        return read_book(z, rows, columns)
    if "content.xml" in names:
        kind = z.read("mimetype").decode("utf-8", "replace").strip() if "mimetype" in names else ""
        if kind == "" or kind.startswith("application/vnd.oasis.opendocument.spreadsheet"):
            return read_ods(z, rows, columns)
        return {"sheets": [], "problem": "Arşiv bir OpenDocument belgesi ama hesap tablosu değil. Tabloyu ODS, XLSX, CSV ya da TXT olarak seçin."}
    return {"sheets": [], "problem": "Arşiv bir Excel (XLSX) ya da OpenDocument (ODS) çalışma kitabı değil. Tabloyu XLSX, ODS, CSV ya da TXT olarak seçin."}


def read(data, rows=ROWS, columns=COLUMNS):
    if len(data) > MAX_BYTES:
        return {"sheets": [], "problem": "Dosya 64 MB'tan büyük; bu kadar büyük tablo okunmuyor."}
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return {"sheets": [], "problem": "Eski Excel biçimi (.xls) okunmuyor. Dosyayı Excel'de XLSX ya da CSV olarak kaydedip seçin."}
    if data[:4] in (b"PK\x03\x04", b"PK\x05\x06"):
        return read_archive(data, rows, columns)
    return read_text(data, rows, columns)


# ── The files ───────────────────────────────────────────────────────────

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def workbook():
    """Two sheets: Koordinatlar (shared, inline and rich strings, a phonetic run, numbers, a boolean, an error, a gap) and Notlar (an absolute target)."""
    shared = (
        f'<sst xmlns="{MAIN}" count="5" uniqueCount="5">'
        "<si><t>Nokta</t></si>"
        "<si><t>Y</t></si>"
        "<si><t>X</t></si>"
        "<si><r><t>Sınır </t></r><r><rPr><b/></rPr><t>taşı</t></r><rPh sb=\"0\" eb=\"1\"><t>yok</t></rPh></si>"
        "<si><t xml:space=\"preserve\"> Ağaç  </t></si>"
        "</sst>"
    )
    sheet1 = (
        f'<worksheet xmlns="{MAIN}"><sheetData>'
        '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" t="s"><v>1</v></c><c r="C1" t="s"><v>2</v></c><c r="D1" t="inlineStr"><is><t>Not</t></is></c></row>'
        '<row r="2"><c r="A2"><v>101</v></c><c r="B2"><v>487000.125</v></c><c r="C2"><v>4420000.0625</v></c><c r="D2" t="s"><v>3</v></c></row>'
        '<row r="3"><c r="A3"><v>102</v></c><c r="B3"><v>1.5E-3</v></c><c r="C3" t="e"><v>#N/A</v></c><c r="D3" t="b"><v>1</v></c></row>'
        '<row r="5"><c r="A5" t="str"><v>P4</v></c><c r="C5" t="b"><v>0</v></c><c t="s"><v>4</v></c></row>'
        '<row><c><v>7</v></c><c r="C6"><v></v></c></row>'
        "</sheetData></worksheet>"
    )
    sheet2 = (
        f'<worksheet xmlns="{MAIN}"><sheetData>'
        '<row r="2"><c r="B2" t="inlineStr"><is><r><t>İl</t></r><r><t>çe</t></r></is></c></row>'
        "</sheetData></worksheet>"
    )
    book = (
        f'<workbook xmlns="{MAIN}" xmlns:r="{RELS}"><sheets>'
        '<sheet name="Koordinatlar" sheetId="1" r:id="rId1"/>'
        '<sheet name="Notlar" sheetId="2" r:id="rId2"/>'
        "</sheets></workbook>"
    )
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/sheet2.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        "</Relationships>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for name, text in [
            ("[Content_Types].xml", CONTENT_TYPES.replace("</Types>", SHEET2_TYPE + "</Types>")),
            ("_rels/.rels", ROOT_RELS),
            ("xl/workbook.xml", book),
            ("xl/_rels/workbook.xml.rels", rels),
            ("xl/sharedStrings.xml", shared),
            ("xl/worksheets/sheet1.xml", sheet1),
            ("xl/worksheets/sheet2.xml", sheet2),
        ]:
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 6, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, text.encode("utf-8"))
    return out.getvalue()


def other_zip():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        info = zipfile.ZipInfo("notlar.txt", date_time=(2026, 10, 6, 0, 0, 0))
        z.writestr(info, b"merhaba")
    return out.getvalue()


def zipped(entries, stored_first=False):
    """A zip of (name, text) entries, dated so the bytes stay the same; the first stored when asked (ODF's mimetype)."""
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for k, (name, text) in enumerate(entries):
            info = zipfile.ZipInfo(name, date_time=(2026, 10, 11, 0, 0, 0))
            info.compress_type = zipfile.ZIP_STORED if stored_first and k == 0 else zipfile.ZIP_DEFLATED
            z.writestr(info, text.encode("utf-8"))
    return out.getvalue()


def dated_workbook(date1904=False):
    """Dates and times: built-in and custom date formats, an elapsed time, a number format, Excel's day 60, a day
    count below 1, a negative one, a string in a date style, Excel's escapes; the 1904 system when asked."""
    styles = (
        f'<styleSheet xmlns="{MAIN}"><numFmts count="5">'
        '<numFmt numFmtId="164" formatCode="dd\\.mm\\.yyyy"/>'
        '<numFmt numFmtId="165" formatCode="[$-41F]d\\ mmmm\\ yyyy;@"/>'
        '<numFmt numFmtId="166" formatCode="0.00&quot; m²&quot;"/>'
        '<numFmt numFmtId="167" formatCode="dd\\.mm\\.yyyy\\ hh:mm:ss"/>'
        '<numFmt numFmtId="168" formatCode="[Red]0.00"/>'
        "</numFmts>"
        '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
        '<fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills>'
        '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
        '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
        '<cellXfs count="9">'
        + "".join(f'<xf numFmtId="{n}" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/>' for n in (0, 14, 164, 165, 20, 46, 166, 167, 168))
        + "</cellXfs>"
        '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
        "</styleSheet>"
    )
    shared = f'<sst xmlns="{MAIN}"><si><t>Tarih</t></si><si><t>Satır_x000D_\nsonu</t></si><si><t>01.10.2024</t></si></sst>'
    rows = [
        '<c r="A1" t="s"><v>0</v></c><c r="B1" t="inlineStr"><is><t>_x005F_x0041_ _x0041_</t></is></c>',
        '<c r="A2" s="1"><v>45566</v></c><c r="B2" s="2"><v>45566</v></c><c r="C2" s="3"><v>45566.75</v></c>',
        '<c r="A3" s="7"><v>45566.5208333333</v></c><c r="B3" s="4"><v>0.75</v></c><c r="C3" s="4"><v>45566.25</v></c>',
        '<c r="A4" s="5"><v>1.5</v></c><c r="B4" s="6"><v>12.5</v></c><c r="C4" s="8"><v>3</v></c>',
        '<c r="A5" s="1"><v>59</v></c><c r="B5" s="1"><v>60</v></c><c r="C5" s="1"><v>61</v></c>',
        '<c r="A6" s="1"><v>0.25</v></c><c r="B6" s="1"><v>-1</v></c><c r="C6" s="1" t="s"><v>2</v></c>',
        '<c r="A7" t="str"><v>a_x000D_b</v></c><c r="B7" t="s"><v>1</v></c><c r="C7" s="2"><v>45566.999999</v></c>',
        '<c r="A8" s="7"><v>45566.999999</v></c><c r="B8" s="1"><v>2958465</v></c><c r="C8" s="1"><v>2958467</v></c>',
    ]
    if date1904:
        rows = ['<c r="A1" s="1"><v>44104</v></c><c r="B1" s="7"><v>1.5</v></c><c r="C1" s="1"><v>0</v></c>']
    sheet = f'<worksheet xmlns="{MAIN}"><sheetData>' + "".join(f'<row r="{i + 1}">{r}</row>' for i, r in enumerate(rows)) + "</sheetData></worksheet>"
    book = (
        f'<workbook xmlns="{MAIN}" xmlns:r="{RELS}">'
        + ('<workbookPr date1904="1"/>' if date1904 else "")
        + '<sheets><sheet name="Tarihler" sheetId="1" r:id="rId1"/></sheets></workbook>'
    )
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
        "</Relationships>"
    )
    return zipped(
        [
            ("[Content_Types].xml", CONTENT_TYPES),
            ("_rels/.rels", ROOT_RELS),
            ("xl/workbook.xml", book),
            ("xl/_rels/workbook.xml.rels", rels),
            ("xl/styles.xml", styles),
            ("xl/sharedStrings.xml", shared),
            ("xl/worksheets/sheet1.xml", sheet),
        ]
    )


CONTENT_TYPES = (
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
    '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
    '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>'
    '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
    "</Types>"
)
SHEET2_TYPE = '<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
ROOT_RELS = (
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
    "</Relationships>"
)

ODS_HEAD = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    f'<office:document-content xmlns:office="{OFFICE}" xmlns:table="{TABLE}" xmlns:text="{TEXTNS}" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'office:version="1.3"><office:body><office:spreadsheet>'
)
ODS_TAIL = "</office:spreadsheet></office:body></office:document-content>"
MANIFEST = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.3">'
    '<manifest:file-entry manifest:full-path="/" manifest:version="1.3" manifest:media-type="{kind}"/>'
    '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
    "</manifest:manifest>"
)


def spreadsheet():
    """Two tables: Noktalar (every value type, a float without its value, string-value, paragraphs with spaces, tabs,
    breaks, spans, a note, an annotation and a drawing; repeated cells and rows, a covered cell with words, header rows
    and a row group, a sheet-long empty tail) and Boş (a table with nothing in it)."""
    cell = lambda attrs, inner="": f"<table:table-cell {attrs}>{inner}</table:table-cell>"
    p = lambda t: f"<text:p>{t}</text:p>"
    rows = [
        "<table:table-header-rows><table:table-row>"
        + cell('office:value-type="string"', p("Ad"))
        + cell('office:value-type="string"', p("Y"))
        + cell('office:value-type="string"', p("X"))
        + cell('office:value-type="string"', p("Tarih"))
        + cell('office:value-type="string"', p("Not"))
        + "</table:table-row></table:table-header-rows>",
        "<table:table-row>"
        + cell('office:value-type="string"', p("P1"))
        + cell('office:value-type="float" office:value="452345.125"', p("452345,13"))
        + cell('office:value-type="float" office:value="4412345.0625"', p("4412345,06"))
        + cell('office:value-type="date" office:date-value="2024-10-01"', p("01.10.2024"))
        + cell('office:value-type="string"', p('Sınır<text:s text:c="2"/>taşı<text:tab/>sol<text:line-break/>alt<text:span> ek</text:span>'
                                              '<text:note><text:note-citation>1</text:note-citation><text:note-body><text:p>dipnot</text:p></text:note-body></text:note> son')
               + p("ikinci satır") + '<office:annotation><text:p>yorum</text:p></office:annotation>')
        + "</table:table-row>",
        "<table:table-row>"
        + cell('office:value-type="string"', p("P2"))
        + cell('office:value-type="percentage" office:value="0.25"', p("%25"))
        + cell('office:value-type="currency" office:currency="TRY" office:value="1340.5"', p("1.340,50 ₺"))
        + cell('office:value-type="date" office:date-value="2024-10-01T12:30:00"', p("01.10.2024 12:30"))
        + cell('office:value-type="time" office:time-value="PT36H05M09.6S"', p("36:05:10"))
        + "</table:table-row>",
        "<table:table-row>"
        + cell('office:value-type="boolean" office:boolean-value="true"', p("DOĞRU"))
        + cell('office:value-type="boolean" office:boolean-value="false"', p("YANLIŞ"))
        + cell('office:value-type="float"', p("12,5 metin"))
        + cell('office:value-type="string" office:string-value="öznitelik"', p("görünen"))
        + cell('table:number-columns-spanned="2" office:value-type="string"', p("birleşik"))
        + '<table:covered-table-cell office:value-type="string"><text:p>gizli</text:p></table:covered-table-cell>'
        + "</table:table-row>",
        '<table:table-row table:number-rows-repeated="2">'
        + cell('table:number-columns-repeated="2" office:value-type="string"', p("tekrar"))
        + cell('table:number-columns-repeated="2"')
        + cell('office:value-type="float" office:value="7"', p("7"))
        + "</table:table-row>",
        "<table:table-row>"
        + cell('office:value-type="string"', "<text:p>\n    girintili  satır\t\n    <text:span>ve  devamı</text:span>\n  </text:p>")
        + "</table:table-row>",
        '<table:table-row-group><table:table-row table:number-rows-repeated="3"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>'
        "<table:table-row>"
        + cell('office:value-type="string"', '<draw:frame><draw:text-box><text:p>çizim</text:p></draw:text-box></draw:frame>' + p("grupta"))
        + "</table:table-row></table:table-row-group>",
        '<table:table-row table:number-rows-repeated="1048566"><table:table-cell table:number-columns-repeated="16384"/></table:table-row>',
    ]
    content = (
        ODS_HEAD
        + '<table:table table:name="Noktalar"><table:table-column table:number-columns-repeated="5"/>'
        + "".join(rows)
        + '</table:table><table:table table:name="Boş"><table:table-row><table:table-cell/></table:table-row></table:table>'
        + ODS_TAIL
    )
    kind = "application/vnd.oasis.opendocument.spreadsheet"
    return zipped([("mimetype", kind), ("META-INF/manifest.xml", MANIFEST.format(kind=kind)), ("content.xml", content)], stored_first=True)


def text_document():
    """An OpenDocument text, not a spreadsheet."""
    kind = "application/vnd.oasis.opendocument.text"
    content = f'<office:document-content xmlns:office="{OFFICE}" xmlns:text="{TEXTNS}"><office:body><office:text><text:p>Merhaba</text:p></office:text></office:body></office:document-content>'
    return zipped([("mimetype", kind), ("META-INF/manifest.xml", MANIFEST.format(kind=kind)), ("content.xml", content)], stored_first=True)


def convert(data, source, ext):
    """LibreOffice's reading of `data` (a file named `source`) written as `ext`: soffice converts it."""
    soffice = shutil.which("soffice")
    if not soffice:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / source
        src.write_bytes(data)
        subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{tmp}/profile", "--convert-to", ext, "--outdir", tmp, str(src)],
                       check=True, capture_output=True, timeout=180)
        return (Path(tmp) / (src.stem + "." + ext)).read_bytes()


def libre(name, source, made):
    """LibreOffice's own writing of a sample (another format), made once and kept."""
    path = FILES / name
    if path.exists():
        return path.read_bytes()
    data = convert(made[source], source, name.rsplit(".", 1)[1])
    if data is None:
        raise SystemExit(f"{name}: LibreOffice (soffice) yok; bu örnek bir kez onunla yazılır.")
    return data


def samples():
    semicolon = (
        "Ada;Parsel;Nitelik;Alan\r\n"
        "101;5;\"Arsa; bahçe\";1340,13\r\n"
        "101;6;\"İki\r\nsatır \"\"tırnak\"\"\";747,00\r\n"
        "\r\n"
        "  102 ;  7 ;Tarla;  \r\n"
    )
    comma = "Nokta,Y,X,Z\n101,487000.125,4420000.500,812.4\n102,487042.5,4419996.75,\n"
    tab = "Ad\tKod\tAçıklama\r\nŞ1\tSN\tŞev üstü ğüşıöç\r\nŞ2\tSN\tİşaret\r\n"
    spaces = "101   487000.125   4420000.500\n102   487042.500   4419996.750\n 103 487047.000 4420024.000\n"
    single = "Birinci satır\n  ikinci, virgüllü ; satır  \nüçüncü\n"
    bar = "a|b|c\n1|2|3\n4||6\n"
    many = "".join(f"{i};{i * 2}\n" for i in range(10_003))
    return [
        ("noktali-virgul.csv", b"\xef\xbb\xbf" + semicolon.encode("utf-8")),
        ("virgul.csv", comma.encode("utf-8")),
        ("sekme-1254.txt", tab.encode("cp1254")),
        ("bosluk.txt", spaces.encode("utf-8")),
        ("tek-sutun.txt", single.encode("utf-8")),
        ("utf16.txt", b"\xff\xfe" + tab.encode("utf-16-le")),
        ("cubuk.csv", bar.encode("utf-8")),
        ("cok-satir.csv", many.encode("utf-8")),
        ("kitap.xlsx", workbook()),
        ("eski.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504),
        ("baska.zip", other_zip()),
        ("tarihler.xlsx", dated_workbook()),
        ("tarihler-1904.xlsx", dated_workbook(date1904=True)),
        ("tablo.ods", spreadsheet()),
        ("metin.odt", text_document()),
    ]


# Samples LibreOffice wrote (kept as made; not rewritten by this script): its XLSX of tablo.ods, its ODS of tarihler.xlsx.
LIBRE = [("libre.xlsx", "tablo.ods"), ("libre.ods", "tarihler.xlsx")]

# Cases read within other limits than Tablo ekle's: (file, rows, columns).
LIMITED = [("tablo.ods", 3, 2), ("kitap.xlsx", 2, 2), ("tarihler.xlsx", 4, 3), ("noktali-virgul.csv", 2, 3)]


def build():
    files = samples()
    made = dict(files)
    kept = [(name, libre(name, source, made)) for name, source in LIBRE]
    made.update(kept)
    cases = [{"file": name, "want": read(data)} for name, data in files + kept]
    for name, rows, columns in LIMITED:
        cases.append({"file": name, "limits": {"rows": rows, "columns": columns}, "want": read(made[name], rows, columns)})
    doc = {
        "format": "kentos.table-file-cases",
        "version": 2,
        "generatedBy": "scripts/fixtures/table_file_cases.py",
        "title": "Tablo dosyaları: CSV, TXT, XLSX ve ODS dosyalarının satırları (ADR 0184 §4, 0223 §1)",
        "processing": PROCESSING,
        "cases": cases,
    }
    return files, kept, doc


def text_of(v):
    return json.dumps(v, ensure_ascii=False, indent=1) + "\n"


# ── Cross-checks ────────────────────────────────────────────────────────


def same(ours, v):
    """Whether the rule's words say what openpyxl read: a date its day (openpyxl gives it its time too), a date and
    time to the second, a time format's value the time of day, an elapsed time the day count, a string with or without
    Excel's escapes (openpyxl reads only _x005F_)."""
    from datetime import timedelta

    hms = lambda secs: f"{secs // 3600:02}:{secs // 60 % 60:02}:{secs % 60:02}"
    if v is None:
        return ours == ""
    if isinstance(v, bool):
        return ours == ("DOĞRU" if v else "YANLIŞ")
    if isinstance(v, datetime):
        if len(ours) == 10:
            return ours == v.date().isoformat()
        r = (v + timedelta(microseconds=500_000)).replace(microsecond=0)
        return ours == (r.isoformat() if "T" in ours else r.time().isoformat())
    if isinstance(v, date):
        return ours == v.isoformat()
    if isinstance(v, time):
        return ours == hms(v.hour * 3600 + v.minute * 60 + v.second + (1 if v.microsecond >= 500_000 else 0))
    if isinstance(v, timedelta):
        try:
            return abs(float(ours) * 86_400 - v.total_seconds()) < 0.001
        except ValueError:
            return ours == hms(math.floor(v.total_seconds() + 0.5))
    if isinstance(v, (int, float)):
        try:
            return float(ours) == float(v)
        except ValueError:
            return False
    return ours == v or ours == unescape(v)


# Cells openpyxl reads otherwise (file, sheet, row, column from 0): why.
OPENPYXL_OTHERWISE = {
    ("tarihler.xlsx", "Tarihler", 5, 1): "eksi gün sayısı: Excel tarih göstermez; kural sayıyı okur, openpyxl 1899'a gider",
    ("tarihler.xlsx", "Tarihler", 7, 2): "9999-12-31'den sonrası: kural sayıyı okur, openpyxl hata (#VALUE!) der",
}

# Cells LibreOffice reads otherwise (file, sheet, row, column from 0): why the rule's words are not LibreOffice's.
LIBRE_OTHERWISE = {
    ("tablo.ods", "Noktalar", 3, 2): "değeri yazılmamış sayı hücresi: kural metnini okur, LibreOffice sayıyı (0)",
    ("tablo.ods", "Noktalar", 2, 4): "36 saatlik süre: LibreOffice XLSX'e günün saati biçimiyle yazar",
    ("libre.ods", "Tarihler", 2, 2): "saat biçimli 45566,25: LibreOffice ODS'ye 1 093 590 saatlik süre yazar, kural süreyi okur",
    ("libre.ods", "Tarihler", 4, 0): "59. gün: LibreOffice var olmayan 29 Şubat 1900'ü saymaz (1900-02-27), Excel ve openpyxl sayar",
    ("libre.ods", "Tarihler", 5, 0): "1'den küçük gün: LibreOffice 30 Aralık 1899'un saati olarak yazar, openpyxl yalnız saat",
    ("libre.ods", "Tarihler", 7, 2): "9999-12-31'den sonrası: LibreOffice 10000 yılını yazar, openpyxl hata der",
}


def compare(name, xlsx, want, otherwise=()):
    """openpyxl's reading of `xlsx` against the rule's sheets: the differing cells."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx), data_only=True)
    out = []
    for k, sheet in enumerate(want.get("sheets", [])):
        ws = wb.worksheets[k]
        merged = {(r - 1, c - 1) for rng in ws.merged_cells.ranges for r in range(rng.min_row, rng.max_row + 1)
                  for c in range(rng.min_col, rng.max_col + 1) if (r, c) != (rng.min_row, rng.min_col)}
        rows = sheet["rows"]
        height = len(rows) if sheet.get("cut") else max(len(rows), ws.max_row)
        for i in range(height):
            width = max(len(rows[i]) if i < len(rows) else 0, ws.max_column)
            for j in range(width):
                if (i, j) in merged or (name, sheet["name"], i, j) in otherwise:
                    continue
                ours = rows[i][j] if i < len(rows) and j < len(rows[i]) else ""
                theirs = ws.cell(i + 1, j + 1).value
                if not same(ours, theirs):
                    out.append(f"{name} › {sheet['name']} {i + 1}:{j + 1}: kural {ours!r}, openpyxl {theirs!r}")
    return out


def cross_check(files):
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("openpyxl yok; çapraz denetim atlandı")
        return []
    out, libre_seen = [], shutil.which("soffice") is not None
    for name, data in files:
        if name.endswith(".xlsx"):
            out += compare(name, data, read(data), OPENPYXL_OTHERWISE)
        elif name.endswith(".ods") and libre_seen:
            out += compare(name, convert(data, name, "xlsx"), read(data), LIBRE_OTHERWISE)
    print(f"cross-checked with openpyxl{' and LibreOffice' if libre_seen else ' (LibreOffice yok: ODS atlandı)'}: {len(out)} fark")
    return out


def main():
    files, kept, doc = build()
    if "--check" in sys.argv[1:]:
        bad = [name for name, data in files if not (FILES / name).exists() or (FILES / name).read_bytes() != data]
        bad += [name for name, _ in kept if not (FILES / name).exists()]
        if not OUT.exists() or OUT.read_text("utf-8") != text_of(doc):
            bad.append(OUT.name)
        for name in bad:
            print(f"{name}: diskteki dosya kurallardan üretilenle aynı değil", file=sys.stderr)
        differ = cross_check(files + kept)
        for d in differ:
            print(d, file=sys.stderr)
        if bad or differ:
            return 1
        print("table file cases match")
        return 0
    FILES.mkdir(parents=True, exist_ok=True)
    for name, data in files + kept:
        (FILES / name).write_bytes(data)
    OUT.write_text(text_of(doc), encoding="utf-8")
    print(f"written: {len(files) + len(kept)} files, {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
