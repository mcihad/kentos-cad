#!/usr/bin/env python3
"""The shared cases of the workbook writer (docs/adr/0223 §2).

    python3 scripts/fixtures/workbook_cases.py           # writes the cases' inputs and expectations
    python3 scripts/fixtures/workbook_cases.py --check   # writes nothing; checks the written files

The writer (`workbook::write`, natively and through WASM: `writeWorkbook`)
writes each case's input; its files are kept in fixtures/table/v1/written
(KENTOS_WRITE_WORKBOOK=1 cargo test -p kentos-formats --test all workbook
rewrites them). This script holds them to the rule, from the ADR on its
own: every written workbook is read back by the table reader's rule
(table_file_cases.py, Python's zipfile and xml.etree) and must give the
words the rule says; every Excel workbook is read by openpyxl (cell types
and values, the names' row bold and frozen, the filter, number formats,
column widths); every OpenDocument spreadsheet is opened by LibreOffice,
written as XLSX and read by openpyxl alike (the filter then LibreOffice's
reading of the database range; its frozen row is not checked: headless
LibreOffice reads no view settings); a CSV is held to this script's own
writing of the rule. Either tool is skipped
when it is not installed.

The rule (a column's kind, then the word):
- number: a decimal number's text (a sign, digits, a point and digits) of at
  most 15 significant digits is a number, its fraction digits shown as the
  column says (0.00), else text; date: a day of the calendar YYYY-MM-DD is a
  date (in an Excel workbook from 1900-03-01, earlier days text), else
  text; boolean: true and false yes and no, else text; auto: a plain number
  (no leading zero, no trailing zero after the point, not -0, at most 15
  significant digits) a number, else text; text: text. An empty word is no
  cell. What became text against its kind is counted in the report.
- A text comes back as written from an Excel workbook (control characters
  as Excel's _xHHHH_); from an OpenDocument spreadsheet its CR LF and CR as
  LF, its control characters but tab and LF left out (counted).
- Sheet names: : \\ / ? * [ ] and control characters become _, apostrophes
  at the ends and spaces round them go, an empty one is SayfaN, at most 31
  characters, History becomes History_, and one already taken (whatever its
  case) gets " (2)", " (3)" … within the 31.
- The names' row is bold and frozen, with a filter over the table; a
  column's width is its widest word (a word's widest line; the name and 3)
  and 2, between 8 and 60 characters.
- CSV: a byte order mark, the names then the rows, fields between
  semicolons, quoted when they hold one, a quote or a line break (a quote
  doubled), lines ending in CR LF.
"""
import io
import json
import shutil
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import table_file_cases as tf  # noqa: E402  (the reader's rule)

ROOT = HERE.parents[1]
DIR = ROOT / "fixtures/table/v1/written"
DIGITS = 15


def significant(t):
    t = t[1:] if t.startswith("-") else t
    whole, _, fraction = t.partition(".")
    return len((whole + fraction.rstrip("0")).lstrip("0"))


def is_decimal(t):
    u = t[1:] if t.startswith("-") else t
    whole, dot, fraction = u.partition(".")
    return whole.isdigit() and whole.isascii() and (not dot or (fraction.isdigit() and fraction.isascii()))


def is_plain(t):
    if not is_decimal(t) or t == "-0":
        return False
    u = t[1:] if t.startswith("-") else t
    whole, _, fraction = u.partition(".")
    return not (len(whole) > 1 and whole.startswith("0")) and not fraction.endswith("0")


def iso_day(t):
    if len(t) != 10 or t[4] != "-" or t[7] != "-" or not (t[:4] + t[5:7] + t[8:]).isdigit():
        return None
    try:
        d = date(int(t[:4]), int(t[5:7]), int(t[8:]))
    except ValueError:
        return None
    return d


def cell(kind, word, excel, counts):
    """('empty'|'text'|'number'|'date'|'boolean', value) as the rule writes the word."""
    if word == "":
        return "empty", None
    if kind == "text":
        return "text", word
    if kind == "number":
        if not is_decimal(word):
            counts["Sayı olmayan değer"] += 1
            return "text", word
        if significant(word) > DIGITS:
            counts["Uzun sayı"] += 1
            return "text", word
        return "number", word
    if kind == "auto":
        return ("number", word) if is_plain(word) and significant(word) <= DIGITS else ("text", word)
    if kind == "date":
        d = iso_day(word)
        if d is None:
            counts["Tarih olmayan değer"] += 1
            return "text", word
        if excel and d < date(1900, 3, 1):
            counts["Eski tarih"] += 1
            return "text", word
        return "date", d
    if kind == "boolean":
        if word in ("true", "false"):
            return "boolean", word == "true"
        counts["Evet/hayır olmayan değer"] += 1
        return "text", word
    raise ValueError(kind)


def sheet_names(names):
    taken, out = set(), []
    for k, n in enumerate(names):
        m = "".join("_" if c in ':\\/?*[]' or ord(c) < 32 or 0x7F <= ord(c) <= 0x9F else c for c in n)
        m = m.strip().strip("'").strip()
        if not m:
            m = f"Sayfa{k + 1}"
        m = m[:31]
        if m.lower() == "history":
            m += "_"
        base, i = m, 1
        while m.lower() in taken:
            i += 1
            tail = f" ({i})"
            m = base[: 31 - len(tail)] + tail
        taken.add(m.lower())
        out.append(m)
    return out


def lines(t):
    """A text's lines as Rust's str::lines gives them: split at LF, a CR before it dropped, no last empty one."""
    parts = t.split("\n")
    if parts and parts[-1] == "":
        parts.pop()
    return [p[:-1] if p.endswith("\r") else p for p in parts]


def width(sheet, j):
    name = len(sheet["columns"][j]["name"]) + 3 if j < len(sheet["columns"]) else 0
    widest = max((max((len(l) for l in lines(r[j])), default=0) for r in sheet["rows"] if j < len(r)), default=0)
    return min(max(max(name, widest) + 2, 8), 60)


def ods_text(t):
    """What an OpenDocument spreadsheet keeps of a text: CR LF and CR as LF, control characters but tab and LF out."""
    t = t.replace("\r\n", "\n").replace("\r", "\n")
    return "".join(c for c in t if c in "\t\n" or not (ord(c) < 32 or 0x7F <= ord(c) <= 0x9F or c in "￾￿"))


def expected(case):
    """The words reading the written file back gives (the reader's rule), and the report's counts."""
    fmt = case["input"]["format"]
    excel = fmt == "xlsx"
    counts = {k: 0 for k in ["Uzun sayı", "Sayı olmayan değer", "Tarih olmayan değer", "Eski tarih", "Evet/hayır olmayan değer"]}
    names = sheet_names([s["name"] for s in case["input"]["sheets"]])
    sheets = []
    for s, name in zip(case["input"]["sheets"], names):
        rows = []
        if s["columns"]:
            rows.append([c["name"] if excel else ods_text(c["name"]) for c in s["columns"]])
        for r in s["rows"]:
            out = []
            for j, w in enumerate(r):
                kind = s["columns"][j]["kind"] if j < len(s["columns"]) else "auto"
                what, v = cell(kind, w, excel, counts)
                out.append({"empty": "", "text": w if excel else ods_text(w), "number": w,
                            "date": v.isoformat() if what == "date" else "", "boolean": "DOĞRU" if v else "YANLIŞ"}[what])
            rows.append(out)
        # The reader's grid: rows end at their last word, the last rows at the last with a word.
        rows = [r[: max((j + 1 for j, w in enumerate(r) if w), default=0)] for r in rows]
        while rows and not rows[-1]:
            rows.pop()
        sheets.append({"name": name, "rows": rows})
    return {"sheets": sheets, "notes": {k: v for k, v in counts.items() if v}}


def csv_text(sheet):
    def field(t):
        return '"' + t.replace('"', '""') + '"' if any(c in t for c in ';"\r\n') else t

    lines = [";".join(field(c["name"]) for c in sheet["columns"])] + [";".join(field(w) for w in r) for r in sheet["rows"]]
    return ("﻿" + "".join(l + "\r\n" for l in lines)).encode("utf-8")


# ── The cases ───────────────────────────────────────────────────────────

ATTRIBUTES = {
    "name": "Parseller",
    "columns": [
        {"name": "Tür", "kind": "text"},
        {"name": "Ada", "kind": "number", "decimals": 0},
        {"name": "Parsel", "kind": "number", "decimals": 0},
        {"name": "Alan (m²)", "kind": "number", "decimals": 2},
        {"name": "Tescil tarihi", "kind": "date"},
        {"name": "Tescilli", "kind": "boolean"},
        {"name": "Malik", "kind": "text"},
        {"name": "Kod", "kind": "auto"},
        {"name": "Not", "kind": "text"},
    ],
    "rows": [
        ["Alan", "101", "5", "1340.13", "2024-10-01", "true", "Ayşe Öztürk", "0.5", "Köşe parseli"],
        ["Alan", "101", "6", "747.00", "1999-02-28", "false", "Çağrı & Şükrü <ortak>", "101/5", "İki\nsatır"],
        ["Alan", "102", "7", "-12.50", "", "", "  Başta ve sonda  ", "007", "Sekme\tile, CR LF\r\nve tek CR\rson"],
        ["Çizgi", "1234567890123456", "1,5", "0.000000000000001234", "1899-12-31", "Evet", "_x0041_ kalır", "1.10", "zil\u0007sesi"],
        ["Nokta", "", "", "", "01.10.2024", "", "", "-0", ""],
        ["Nokta", "9007199254740991", "8", "452345.123", "2023-02-29", "true", "Son", "-3.25", "  "],
        [],
    ],
}

NAMED = [
    {"name": "Ada/Parsel", "columns": [{"name": "A", "kind": "text"}], "rows": [["bir"]]},
    {"name": "ada_parsel", "columns": [{"name": "B", "kind": "auto"}], "rows": [["2"], ["3.5"]]},
    {"name": "'Yol'", "columns": [], "rows": []},
    {"name": "", "columns": [{"name": "C", "kind": "text"}], "rows": [["üç"]]},
    {"name": "Çok uzun bir katman adı ki otuz bir karakteri aşar", "columns": [{"name": "D", "kind": "number", "decimals": 3}], "rows": [["1.5"]]},
    {"name": "History", "columns": [{"name": "E", "kind": "date"}], "rows": [["2000-01-01"]]},
]

CASES = [
    ("oznitelik", {"format": "xlsx", "sheets": [ATTRIBUTES]}),
    ("oznitelik", {"format": "ods", "sheets": [ATTRIBUTES]}),
    ("oznitelik", {"format": "csv", "sheets": [ATTRIBUTES]}),
    ("adlar", {"format": "xlsx", "sheets": NAMED}),
    ("adlar", {"format": "ods", "sheets": NAMED}),
]


def cases():
    out = []
    for name, data in CASES:
        case = {"name": f"{name}.{data['format']}", "input": data}
        if data["format"] != "csv":
            case["expect"] = expected(case)
        out.append(case)
    return out


def text_of(v):
    return json.dumps(v, ensure_ascii=False, indent=1) + "\n"


# ── Checks of the written files ─────────────────────────────────────────


def read_words(data):
    got = tf.read(data, tf.PROCESSING["rows"], tf.PROCESSING["columns"])
    return [{"name": s["name"], "rows": s["rows"]} for s in got["sheets"]]


def check_openpyxl(case, data, through):
    """openpyxl's reading of an Excel workbook (`through` names how it was made) against the rule."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data))
    out = []
    names = sheet_names([s["name"] for s in case["input"]["sheets"]])
    if wb.sheetnames != names:
        out.append(f"{case['name']} ({through}): sayfa adları {wb.sheetnames}, kural {names}")
        return out
    for s, ws in zip(case["input"]["sheets"], wb.worksheets):
        where = f"{case['name']} ({through}) › {ws.title}"
        if s["columns"]:
            # Headless LibreOffice reads no view settings: an ODS's frozen row is not seen through it.
            if through == "yazıcı" and ws.freeze_panes != "A2":
                out.append(f"{where}: dondurulan {ws.freeze_panes}, kural A2")
            width_ = max([len(r) for r in s["rows"]] + [len(s["columns"]), 1])
            want_ref = f"A1:{column(width_ - 1)}{len(s['rows']) + 1}"
            if ws.auto_filter.ref != want_ref:
                out.append(f"{where}: süzgeç {ws.auto_filter.ref}, kural {want_ref}")
            for j, c in enumerate(s["columns"]):
                head = ws.cell(1, j + 1)
                if head.value != c["name"] or not head.font.b:
                    out.append(f"{where}: başlık {j + 1} {head.value!r} kalın={head.font.b}")
                if through == "yazıcı":
                    w = ws.column_dimensions[column(j)].width
                    if abs(w - width(s, j)) > 1e-9:
                        out.append(f"{where}: {column(j)} genişliği {w}, kural {width(s, j)}")
        counts = {k: 0 for k in ["Uzun sayı", "Sayı olmayan değer", "Tarih olmayan değer", "Eski tarih", "Evet/hayır olmayan değer"]}
        top = 2 if s["columns"] else 1
        for i, r in enumerate(s["rows"]):
            for j, w in enumerate(r):
                kind = s["columns"][j]["kind"] if j < len(s["columns"]) else "auto"
                what, v = cell(kind, w, through == "yazıcı" or case["input"]["format"] == "xlsx", counts)
                got = ws.cell(i + top, j + 1)
                gv = got.value
                if (case["name"], i + top, j + 1) in LIBRE_OTHERWISE and through == "LibreOffice":
                    continue
                ok = {
                    "empty": gv is None,
                    "text": isinstance(gv, str) and (tf.unescape(gv) == w or gv == ods_text(w) or tf.unescape(gv) == ods_text(w)),
                    "number": isinstance(gv, (int, float)) and not isinstance(gv, bool) and float(gv) == float(w),
                    "date": isinstance(gv, datetime) and gv.date() == v,
                    "boolean": gv is v,
                }[what]
                if not ok:
                    out.append(f"{where} {i + top}:{j + 1}: {what} {w!r}, openpyxl {gv!r}")
                if what == "number" and through == "yazıcı" and kind == "number" and s["columns"][j].get("decimals") is not None:
                    d = s["columns"][j]["decimals"]
                    want = "0" if d == 0 else "0." + "0" * d
                    if got.number_format != want:
                        out.append(f"{where} {i + top}:{j + 1}: biçim {got.number_format!r}, kural {want!r}")
    return out


# Cells LibreOffice writes otherwise to XLSX (case, row, column from 1): why.
LIBRE_OTHERWISE = {
    ("oznitelik.ods", 5, 5): "1899-12-31: LibreOffice XLSX'e kendi gün sayısını (1) yazar, Excel'in 1'i 1900-01-01'dir",
}


def column(j):
    out = ""
    j += 1
    while j:
        j, r = divmod(j - 1, 26)
        out = chr(65 + r) + out
    return out


def check():
    problems, made = [], cases()
    doc_path = DIR / "cases.json"
    if not doc_path.exists() or doc_path.read_text("utf-8") != text_of({"format": "kentos.workbook-cases", "version": 1, "cases": made}):
        problems.append("cases.json: diskteki dosya kurallardan üretilenle aynı değil")
    try:
        import openpyxl  # noqa: F401

        have_openpyxl = True
    except ImportError:
        have_openpyxl = False
        print("openpyxl yok; XLSX denetimi atlandı")
    soffice = shutil.which("soffice")
    if not soffice:
        print("LibreOffice yok; ODS denetimi atlandı")
    for case in made:
        path = DIR / case["name"]
        if not path.exists():
            problems.append(f"{case['name']}: yazılmış dosya yok (KENTOS_WRITE_WORKBOOK=1 cargo test -p kentos-formats --test all workbook)")
            continue
        data = path.read_bytes()
        fmt = case["input"]["format"]
        if fmt == "csv":
            if data != csv_text(case["input"]["sheets"][0]):
                problems.append(f"{case['name']}: CSV kuralın yazdığından farklı")
            continue
        if zipfile.ZipFile(io.BytesIO(data)).testzip() is not None:
            problems.append(f"{case['name']}: ZIP bozuk")
        if fmt == "ods":
            z = zipfile.ZipFile(io.BytesIO(data))
            first = z.infolist()[0]
            if first.filename != "mimetype" or first.compress_type != zipfile.ZIP_STORED or z.read("mimetype") != b"application/vnd.oasis.opendocument.spreadsheet":
                problems.append(f"{case['name']}: mimetype ilk ve sıkıştırmasız değil")
        if read_words(data) != case["expect"]["sheets"]:
            problems.append(f"{case['name']}: geri okunan sözcükler kuralınkinden farklı: {read_words(data)}")
        if have_openpyxl and fmt == "xlsx":
            problems += check_openpyxl(case, data, "yazıcı")
        if have_openpyxl and fmt == "ods" and soffice:
            xlsx = tf.convert(data, case["name"], "xlsx")
            problems += check_openpyxl(case, xlsx, "LibreOffice")
    return problems


def main():
    made = cases()
    if "--check" in sys.argv[1:]:
        problems = check()
        for p in problems:
            print(p, file=sys.stderr)
        if problems:
            return 1
        print(f"workbook cases match ({len(made)} cases)")
        return 0
    DIR.mkdir(parents=True, exist_ok=True)
    (DIR / "cases.json").write_text(text_of({"format": "kentos.workbook-cases", "version": 1, "cases": made}), encoding="utf-8")
    print(f"written: {DIR.relative_to(ROOT)}/cases.json ({len(made)} cases)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
